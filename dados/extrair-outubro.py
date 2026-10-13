# -*- coding: utf-8 -*-
"""Extrai a contagem de OUTUBRO/2026 do Bar do Zeca - Norte Shopping.

Mesmo desenho do extrair.py de agosto: cada item e conferido contra o total da
propria planilha, e cada bloco contra a linha TOTAL do bloco. O que mudou de
agosto para ca:

  1. o SETOR nao vem mais da aba. A migracao 0002 juntou os 6 setores de agosto
     em 4 (Bar, Cozinha, Limpeza, Descartaveis), e a planilha continua com o
     layout antigo. Quem decide o setor aqui e o par (aba, bloco) — ver SETOR_DE.
  2. cada item carrega `aba` e `bloco` de origem. Sem isso o normalizar nao tem
     como desempatar os 16 nomes que existem em dois produtos diferentes
     (ALECRIM x ALECRIM (HORTIFRUTI), BACON x BACON (SUINOS) x BACON (FEIJOADA)),
     e errar isso troca o custo do item e corrompe o CMV.
  3. o resumo da propria planilha (Geral A74:E95) e lido e comparado com a soma
     por categoria. E a conferencia mais forte que existe no arquivo: e o numero
     que a casa usa, calculado por formula do Excel sobre as mesmas celulas.

Escreve dados/contagem-bruta-outubro.json.
"""
import json, os, re
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(AQUI, 'CONTAGEM_BDZ_NORTE_SHOPPING_01_OUTUBRO_2026.xlsx')
DEST = os.path.join(AQUI, 'contagem-bruta-outubro.json')

# Tolerancia das conferencias. A planilha guarda float binario (os totais saem
# como 507.79400000000004), entao a comparacao e por valor, nao por igualdade.
TOLERANCIA = 0.02

wb = openpyxl.load_workbook(SRC, data_only=True)


def num(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('R$', '').replace(' ', '')
    if ',' in s and '.' in s: s = s.replace('.', '').replace(',', '.')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(s)
    except ValueError: return 0.0


def txt(v):
    if v is None: return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def eh_total(s):
    u = txt(s).upper()
    return u.startswith('TOTAL') or u.startswith('VALOR TOTAL') or u.startswith('VALOR FINAL') or u.endswith('TOTAL')


# --------------------------------------------------------------------- setor --
# De qual setor ATUAL do sistema cada bloco da planilha fala. A planilha e a
# mesma de agosto, de quando existiam 6 setores; o sistema tem 4. Nos dois casos
# em que o bloco do `Geral` sai da Cozinha (DESCARTAVEIS e MATERIAL DE LIMPEZA)
# a decisao e pelo nome do bloco, nao pela aba.
SETOR_POR_BLOCO_GERAL = {
    'DESCARTÁVEIS': 'Descartáveis',
    'MATERIAL DE LIMPEZA': 'Limpeza',
}

SETOR_POR_ABA = {
    'Geral': 'Cozinha',
    'Hortifruti': 'Cozinha',
    'Massas': 'Cozinha',
    'Porcionados': 'Cozinha',
    'Molhos': 'Cozinha',
    'Bebidas Álcoolicas': 'Bar',
    'Bebidas não Álcoolicas': 'Bar',
}


def setor_de(aba, bloco):
    if aba == 'Geral' and bloco in SETOR_POR_BLOCO_GERAL:
        return SETOR_POR_BLOCO_GERAL[bloco]
    return SETOR_POR_ABA[aba]


itens = []        # registros finais, uma linha por celula de produto
avisos = []       # divergencia item x item (quantidade * valor != TOTAL da planilha)
conferencia = []  # total do bloco: planilha x calculado
achados = []      # qualquer coisa que a planilha trouxe fora do esperado


def add(aba, bloco, categoria, nome, qtd, unidade, custo, total_planilha, origem):
    """Registra uma linha e confere o produto dela contra o TOTAL da planilha.

    Nao filtra linha de titulo nem nome vazio alem do minimo (nome em branco,
    '0' e linha de total): o descarte com justificativa e trabalho do
    normalizar-outubro.py, que precisa listar o que jogou fora.
    """
    nome = txt(nome)
    if not nome or nome == '0' or eh_total(nome): return
    # `total_item` e a conta refeita do zero, com o mesmo arredondamento da
    # coluna gerada contagem_itens.total (round(quantidade * custo, 4)). E ele,
    # nao o numero da planilha, que o seed vai ter de reproduzir no banco.
    q, c = round(qtd, 4), round(custo, 6)
    calc = round(q * c, 4)
    tp = round(num(total_planilha), 4)
    if abs(calc - tp) > TOLERANCIA:
        avisos.append({'origem': origem, 'aba': aba, 'bloco': bloco, 'produto': nome,
                       'quantidade': q, 'custo': c,
                       'total_planilha': tp, 'total_calculado': calc,
                       'diferenca': round(calc - tp, 4)})
    itens.append({'aba': aba, 'bloco': bloco, 'setor': setor_de(aba, bloco),
                  'categoria': categoria, 'produto': nome,
                  'quantidade': q, 'unidade': txt(unidade).upper() or 'UND',
                  'custo_unitario': c, 'total': tp, 'total_item': calc,
                  'origem': origem})


# ---------- Aba Geral: 4 blocos de 7 colunas, cada um com varias categorias ----------
# Ultima linha util de cada bloco. Conferida contra a planilha de outubro: a
# linha LIMITE e sempre a do TOTAL do ultimo bloco da coluna, e GUARDA_GERAL
# abaixo prova que nao sobrou produto fora do alcance.
LIMITE = {1: 71, 8: 78, 15: 71, 22: 144}
# No bloco 1 as linhas 74..95 sao o resumo da casa (MERCADORIAS EM ESTOQUE), nao
# produto. Elas sao lidas em separado, como conferencia.
RESUMO = (74, 95)

ws = wb['Geral']
for c0 in (1, 8, 15, 22):
    categoria = None
    soma = 0.0
    total_visto = False   # o bloco corrente ja encontrou a propria linha de TOTAL?
    for r in range(1, LIMITE[c0] + 1):
        a = txt(ws.cell(r, c0).value)
        b = txt(ws.cell(r, c0 + 1).value)
        if a and b.upper() == 'ESTOQUE':          # cabecalho de categoria
            # Bloco que terminou sem linha de TOTAL fica registrado assim mesmo:
            # e o unico jeito de o relatorio dizer que ninguem conferiu aquele
            # pedaco contra a planilha, em vez de a ausencia passar batida.
            if categoria and not total_visto:
                conferencia.append({'bloco': f'Geral/{categoria}', 'calculado': round(soma, 4),
                                    'planilha': None, 'origem': None})
            categoria, soma, total_visto = a, 0.0, False
            continue
        if not categoria: continue
        if eh_total(a) or (not a and num(ws.cell(r, c0 + 4).value)):   # linha de total (com ou sem rotulo)
            tot = num(ws.cell(r, c0 + 4).value)
            if tot or not total_visto:
                conferencia.append({'bloco': f'Geral/{categoria}', 'planilha': round(tot, 4),
                                    'calculado': round(soma, 4),
                                    'origem': f'Geral!{ws.cell(r, c0 + 4).coordinate}'})
                total_visto, soma = True, 0.0
            continue
        if not a: continue
        q, un, vl, tt = (num(ws.cell(r, c0 + 1).value), ws.cell(r, c0 + 2).value,
                         num(ws.cell(r, c0 + 3).value), ws.cell(r, c0 + 4).value)
        add('Geral', categoria, categoria, a, q, un, vl, tt, f'Geral!{ws.cell(r, c0).coordinate}')
        soma += num(tt)

# GUARDA_GERAL: nada de produto pode viver abaixo de LIMITE. Se a cliente
# acrescentar uma categoria no fim de uma coluna, o extrator passaria batido e o
# total fecharia errado "sem motivo" — por isso a varredura reclama em voz alta.
for c0 in (1, 8, 15, 22):
    for r in range(LIMITE[c0] + 1, ws.max_row + 1):
        if c0 == 1 and RESUMO[0] <= r <= RESUMO[1]:
            continue
        nome = txt(ws.cell(r, c0).value)
        if nome and (num(ws.cell(r, c0 + 1).value) or num(ws.cell(r, c0 + 3).value)):
            achados.append(f'Geral!{ws.cell(r, c0).coordinate}: {nome!r} abaixo do '
                           f'limite da coluna (linha {LIMITE[c0]}) — fora da extracao')


# ---------- Abas com layout simples: nome | q1 | q2 | valor | total ----------
def aba_simples(nome_aba, cat_padrao, col_nome=1, dupla=True, rotulo_total_da_aba=None):
    """Le uma aba de coluna unica.

    `dupla` liga a regra das abas Massas / Porcionados / Molhos: ha duas colunas
    de contagem (PORCAO/PORCOES e UND/ESTOQUE-KG) que se somam, e a unidade e
    UND quando a primeira tem numero, senao KG. A regra e de agosto e esta
    conferida contra os totais da planilha de outubro (ver conferencia por
    bloco): se a soma de um bloco com `dupla` fechar, a regra continua valendo.

    `rotulo_total_da_aba` e o rotulo da linha que fecha a ABA inteira, nao um
    bloco — `TOTAL PORCIONADOS:` na linha 117 de Porcionados vem DEPOIS do
    `TOTAL` do ultimo bloco e soma os seis blocos. Sem distinguir os dois, essa
    linha seria cobrada do bloco PROCESSADOS e apareceria como divergencia de
    R$ 4.646,13 que nao existe.
    """
    ws = wb[nome_aba]
    categoria = cat_padrao
    soma = 0.0
    soma_aba = 0.0
    total_visto = False
    for r in range(1, ws.max_row + 1):
        a = txt(ws.cell(r, col_nome).value)
        b = txt(ws.cell(r, col_nome + 1).value)
        if not a: continue
        if b.upper() in ('ESTOQUE', 'PORÇAO', 'PORÇÕES', 'PORCOES', 'QUANTIDADE') or \
           (b.upper() == '' and txt(ws.cell(r, col_nome + 3).value).upper().startswith('VALOR')):
            if (categoria != cat_padrao or soma) and not total_visto:
                conferencia.append({'bloco': f'{nome_aba}/{categoria}', 'calculado': round(soma, 4),
                                    'planilha': None, 'origem': None})
            categoria, soma, total_visto = a, 0.0, False
            continue
        if eh_total(a):
            tot = num(ws.cell(r, col_nome + 4).value)
            celula = f'{nome_aba}!{ws.cell(r, col_nome + 4).coordinate}'
            if rotulo_total_da_aba and a.upper().rstrip(':').strip() == rotulo_total_da_aba:
                conferencia.append({'bloco': f'{nome_aba} (aba inteira)', 'planilha': round(tot, 4),
                                    'calculado': round(soma_aba, 4), 'origem': celula})
            else:
                conferencia.append({'bloco': f'{nome_aba}/{categoria}', 'planilha': round(tot, 4),
                                    'calculado': round(soma, 4), 'origem': celula})
                total_visto, soma = True, 0.0
            continue
        tt = ws.cell(r, col_nome + 4).value
        vl = num(ws.cell(r, col_nome + 3).value)
        if dupla:   # duas colunas de contagem (porcao + unidade/kg) que se somam
            q = num(ws.cell(r, col_nome + 1).value) + num(ws.cell(r, col_nome + 2).value)
            un = 'UND' if num(ws.cell(r, col_nome + 1).value) else 'KG'
        else:
            q = num(ws.cell(r, col_nome + 1).value)
            un = ws.cell(r, col_nome + 2).value
        add(nome_aba, categoria, categoria, a, q, un, vl, tt, f'{nome_aba}!A{r}')
        soma += num(tt)
        soma_aba += num(tt)


aba_simples('Hortifruti', 'HORTIFRUTI', dupla=False)
aba_simples('Massas', 'MASSAS')
aba_simples('Porcionados', 'CARNES', rotulo_total_da_aba='TOTAL PORCIONADOS')
aba_simples('Molhos', 'MOLHOS')
# Só a coluna A desta aba: a partir da coluna 12 existe um bloco "BEBIDAS PEDRO
# PACHECO" que e lista de pedido de fornecedor, tudo zerado, e nao e contagem.
aba_simples('Bebidas Álcoolicas', 'BEBIDAS ALCOOLICAS', dupla=False)
aba_simples('Bebidas não Álcoolicas', 'BEBIDAS NAO ALCOOLICAS', dupla=False)

# A coluna 12+ da aba de alcoolicas fica registrada como achado, com os valores,
# para o relatorio poder afirmar que ela foi vista e deixada de fora de proposito.
wsb = wb['Bebidas Álcoolicas']
pedido_fornecedor = []
for r in range(1, wsb.max_row + 1):
    rotulo = txt(wsb.cell(r, 12).value)
    if not rotulo: continue
    pedido_fornecedor.append({'origem': f'Bebidas Álcoolicas!{wsb.cell(r, 12).coordinate}',
                              'rotulo': rotulo,
                              'valor': num(wsb.cell(r, 15).value)})

# A aba Perda tambem nao e contagem de estoque. Vai inteira para o relatorio.
wsp = wb['Perda']
perdas = []
for r in range(1, wsp.max_row + 1):
    linha = [txt(wsp.cell(r, c).value) for c in range(1, wsp.max_column + 1)]
    if any(linha): perdas.append({'linha': r, 'celulas': linha})


# ------------------------------------------------------------------- resumo --
# O quadro MERCADORIAS EM ESTOQUE (Geral A74:E95) e o numero da casa. Cada linha
# e <rotulo do resumo> = <total>, e VALOR FINAL e o total da contagem.
# O rotulo do resumo nao e igual ao nome do bloco (PROTEINA CLIENTE x PROTEINAS
# CLIENTE, COMIDA FUNCIONARIO x PROTEINAS FUNCIONARIO, LATICINIOS x LATICINIOS),
# por isso o de-para e explicito: adivinhar aqui seria conferir contra si mesmo.
RESUMO_PARA_BLOCOS = {
    'BEBIDAS ÁLCOOLICAS': ['Bebidas Álcoolicas/BEBIDAS ALCOOLICAS'],
    'BEBIDAS NÃO ÁLCOOLICAS': ['Bebidas não Álcoolicas/BEBIDAS NÃO ALCOÓLICAS (ESTOQUE)'],
    'CONGELADOS': ['Geral/CONGELADOS'],
    'COMIDA FUNCIONÁRIO': ['Geral/PROTEÍNAS FUNCIONÁRIO'],
    'DESCARTÁVEIS': ['Geral/DESCARTÁVEIS'],
    'EMBUTIDOS': ['Geral/EMBUTIDOS'],
    'HORTIFRUTI': ['Hortifruti/PRODUTOS'],
    'LATICINIOS': ['Geral/LATICÍNIOS'],
    'MATERIAL DE LIMPEZA': ['Geral/MATERIAL DE LIMPEZA'],
    'MASSAS': ['Massas/MASSAS'],
    'MERCEARIA': ['Geral/MERCEARIA'],
    'MOLHOS': ['Molhos/MOLHOS'],
    # PORCIONADOS no resumo e a aba inteira: os 6 blocos dela somados.
    'PORCIONADOS': ['Porcionados/CARNES', 'Porcionados/APARAS',
                    'Porcionados/PEIXES E FRUTOS DO MAR', 'Porcionados/SUÍNOS',
                    'Porcionados/FEIJOADA', 'Porcionados/PROCESSADOS'],
    'PROTEINA CLIENTE': ['Geral/PROTEÍNAS CLIENTE'],
    'PEIXES E FRUTOS DO MAR': ['Geral/PEIXES E FRUTOS DO MAR'],
    'SORVETES': ['Geral/SORVETE'],
    'TEMPEROS E AFINS': ['Geral/TEMPEROS E AFINS'],
}

resumo = []
valor_final = None
for r in range(RESUMO[0], RESUMO[1] + 1):
    rotulo = txt(ws.cell(r, 1).value).rstrip(':').strip()
    if not rotulo or rotulo in ('NORTE SHOPPING', 'MERCADORIAS EM ESTOQUE'): continue
    valor = round(num(ws.cell(r, 2).value), 4)
    if rotulo == 'VALOR FINAL':
        valor_final = valor
        continue
    if rotulo == 'PERDA DE PRATO':
        achados.append(f'Geral!B{r}: PERDA DE PRATO = {valor} (fora da contagem de estoque)')
        continue
    blocos = RESUMO_PARA_BLOCOS.get(rotulo)
    if blocos is None:
        achados.append(f'Geral!A{r}: rotulo de resumo {rotulo!r} sem bloco correspondente')
        continue
    calculado = round(sum(i['total'] for i in itens
                          if f"{i['aba']}/{i['bloco']}" in blocos), 4)
    resumo.append({'rotulo': rotulo, 'origem': f'Geral!B{r}', 'blocos': blocos,
                   'planilha': valor, 'calculado': calculado,
                   'diferenca': round(calculado - valor, 4)})

soma_resumo = round(sum(x['planilha'] for x in resumo), 4)
# Dois totais, de proposito. `total_planilha` soma a coluna TOTAL (R$) como ela
# esta na planilha; `total_itens` refaz quantidade x valor com o arredondamento
# do banco. Sao o mesmo numero quando a planilha esta consistente — e e o
# segundo que o seed tem de reproduzir.
total_planilha = round(sum(i['total'] for i in itens), 4)
total_itens = round(sum(i['total_item'] for i in itens), 4)

out = {
    'restaurante': 'Bar do Zeca — Norte Shopping',
    'competencia': '2026-10',
    'origem': os.path.basename(SRC),
    'total_planilha': total_planilha,
    'total_itens_recalculado': total_itens,
    'total_planilha_valor_final': valor_final,
    'soma_das_linhas_do_resumo': soma_resumo,
    'itens': itens,
    'avisos': avisos,
    'conferencia': conferencia,
    'resumo_planilha': resumo,
    'pedido_fornecedor_fora_da_contagem': pedido_fornecedor,
    'aba_perda_fora_da_contagem': perdas,
    'achados': achados,
}
with open(DEST, 'w') as fh:
    json.dump(out, fh, ensure_ascii=False, indent=1)

print('arquivo             :', os.path.basename(SRC))
print('itens lidos         :', len(itens))
print('soma da coluna TOTAL: %.4f' % total_planilha)
print('soma recalculada    : %.4f   (quantidade x valor, como o banco soma)' % total_itens)
print('VALOR FINAL planilha: %.4f   (Geral!B93)    diferenca: %+.6f'
      % (valor_final, total_itens - valor_final))
print('soma do resumo      : %.4f   (as 17 linhas) diferenca: %+.6f'
      % (soma_resumo, soma_resumo - valor_final))

print('\n--- divergencias item x item (quantidade * valor != TOTAL):', len(avisos))
for a in avisos: print('   ', a)

print('\n--- conferencia por bloco (TOTAL do bloco x soma das linhas) ---')
ruins = 0
for c in conferencia:
    p, k = c['planilha'], c['calculado']
    if p is None:
        marca = '(bloco sem linha de TOTAL na planilha)'
    elif abs(p - k) < TOLERANCIA:
        marca = 'OK'
    else:
        marca = '<<< DIVERGE'
        ruins += 1
    print('   %-52s planilha=%-14s calculado=%-14s %s' % (c['bloco'], p, k, marca))
print('   blocos que divergem:', ruins)

print('\n--- conferencia contra o resumo da planilha (Geral A74:E95) ---')
ruins_resumo = 0
for x in resumo:
    ok = abs(x['diferenca']) < TOLERANCIA
    if not ok: ruins_resumo += 1
    print('   %-26s %-10s planilha=%-14s calculado=%-14s %s'
          % (x['rotulo'], x['origem'], x['planilha'], x['calculado'], 'OK' if ok else '<<< DIVERGE'))
print('   linhas de resumo que divergem:', ruins_resumo)

print('\n--- fora da contagem ---')
print('   bloco de pedido de fornecedor (coluna 12 das alcoolicas):',
      len(pedido_fornecedor), 'linhas, soma %.4f' % sum(x['valor'] for x in pedido_fornecedor))
print('   aba Perda:', len(perdas), 'linhas com conteudo')
for p in perdas: print('      r%-3d %s' % (p['linha'], p['celulas']))

print('\n--- achados ---')
for a in achados: print('   ', a)
if not achados: print('    nenhum')

print('\nescrito:', DEST)
