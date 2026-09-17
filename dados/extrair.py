# -*- coding: utf-8 -*-
"""Extrai a contagem de Agosto do Bar do Zeca - Norte Shopping.
Cada item e conferido contra o total da propria planilha."""
import json, re, unicodedata
import openpyxl

SRC = '/root/.claude/uploads/ea7223e1-0cfd-5570-ae56-c7244317aec2/6cb3361a-Contagem_BDZ_NS_08_Agosto_2026_okk.xlsx'
wb = openpyxl.load_workbook(SRC, data_only=True)

def num(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace('R$', '').replace(' ', '')
    if ',' in s and '.' in s: s = s.replace('.', '').replace(',', '.')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(s)
    except ValueError: return 0.0

def txt(v):
    if v is None: return ''
    return re.sub(r'\s+', ' ', str(v)).strip()

def eh_total(s):
    u = txt(s).upper()
    return u.startswith('TOTAL') or u.startswith('VALOR TOTAL') or u.startswith('VALOR FINAL') or u.endswith('TOTAL')

itens = []          # registros finais
avisos = []         # divergencias encontradas
conferencia = []    # total planilha x total calculado, por bloco

def add(setor, categoria, nome, qtd, unidade, custo, total_planilha, origem):
    nome = txt(nome)
    if not nome or nome == '0' or eh_total(nome): return
    calc = round(qtd * custo, 4)
    tp = round(num(total_planilha), 4)
    if abs(calc - tp) > 0.02:
        avisos.append({'origem': origem, 'produto': nome, 'qtd': qtd, 'custo': custo,
                       'total_planilha': tp, 'total_calculado': calc})
    itens.append({'setor': setor, 'categoria': categoria, 'produto': nome,
                  'quantidade': round(qtd, 4), 'unidade': txt(unidade).upper() or 'UND',
                  'custo_unitario': round(custo, 6), 'total': tp, 'origem': origem})

# ---------- Aba Geral: 4 blocos de 7 colunas, cada um com varias categorias ----------
ws = wb['Geral']
LIMITE = {1: 71, 8: 78, 15: 71, 22: 144}   # ultima linha util de cada bloco (row 74+ do bloco 1 e resumo)
for c0 in (1, 8, 15, 22):
    categoria = None
    soma = 0.0
    for r in range(1, LIMITE[c0] + 1):
        a = txt(ws.cell(r, c0).value)
        b = txt(ws.cell(r, c0 + 1).value)
        if a and b.upper() == 'ESTOQUE':          # cabecalho de categoria
            if categoria: conferencia.append({'bloco': f'Geral/{categoria}', 'calculado': round(soma, 4)})
            categoria, soma = a, 0.0
            continue
        if not categoria: continue
        if eh_total(a) or (not a and num(ws.cell(r, c0 + 4).value)):   # linha de total (com ou sem rotulo)
            tot = num(ws.cell(r, c0 + 4).value)
            if tot: conferencia.append({'bloco': f'Geral/{categoria}', 'planilha': round(tot, 4), 'calculado': round(soma, 4)})
            continue
        if not a: continue
        q, un, vl, tt = (num(ws.cell(r, c0 + 1).value), ws.cell(r, c0 + 2).value,
                         num(ws.cell(r, c0 + 3).value), ws.cell(r, c0 + 4).value)
        add('Estoque Geral', categoria, a, q, un, vl, tt, f'Geral!{ws.cell(r, c0).coordinate}')
        soma += num(tt)

# ---------- Abas com layout simples: nome | q1 | q2 | valor | total ----------
def aba_simples(nome_aba, setor, cat_padrao, col_nome=1, dupla=True):
    ws = wb[nome_aba]
    categoria = cat_padrao
    soma = 0.0
    for r in range(1, ws.max_row + 1):
        a = txt(ws.cell(r, col_nome).value)
        b = txt(ws.cell(r, col_nome + 1).value)
        if not a: continue
        if b.upper() in ('ESTOQUE', 'PORÇAO', 'PORÇÕES', 'PORCOES', 'QUANTIDADE') or \
           (b.upper() == '' and txt(ws.cell(r, col_nome + 3).value).upper().startswith('VALOR')):
            if categoria != cat_padrao or soma:
                conferencia.append({'bloco': f'{nome_aba}/{categoria}', 'calculado': round(soma, 4)})
            categoria, soma = a, 0.0
            continue
        if eh_total(a):
            tot = num(ws.cell(r, col_nome + 4).value)
            conferencia.append({'bloco': f'{nome_aba}/{categoria}', 'planilha': round(tot, 4), 'calculado': round(soma, 4)})
            soma = 0.0
            continue
        tt = ws.cell(r, col_nome + 4).value
        vl = num(ws.cell(r, col_nome + 3).value)
        if dupla:   # duas colunas de contagem (porcao + unidade/kg) que se somam
            q = num(ws.cell(r, col_nome + 1).value) + num(ws.cell(r, col_nome + 2).value)
            un = 'UND' if num(ws.cell(r, col_nome + 1).value) else 'KG'
        else:
            q = num(ws.cell(r, col_nome + 1).value)
            un = ws.cell(r, col_nome + 2).value
        add(setor, categoria, a, q, un, vl, tt, f'{nome_aba}!A{r}')
        soma += num(tt)

aba_simples('Hortifruti', 'Hortifruti', 'HORTIFRUTI', dupla=False)
aba_simples('Massas', 'Massas e Panificacao', 'MASSAS')
aba_simples('Porcionados', 'Porcionados', 'CARNES')
aba_simples('Molhos', 'Molhos e Caldos', 'MOLHOS')
aba_simples('Bebidas Álcoolicas', 'Bar', 'BEBIDAS ALCOOLICAS', dupla=False)
aba_simples('Bebidas não Álcoolicas', 'Bar', 'BEBIDAS NAO ALCOOLICAS', dupla=False)

out = {'restaurante': 'Bar do Zeca - Norte Shopping', 'competencia': '2026-08',
       'itens': itens, 'avisos': avisos, 'conferencia': conferencia}
dest = '/tmp/claude-0/-home-user/ea7223e1-0cfd-5570-ae56-c7244317aec2/scratchpad/contagem.json'
json.dump(out, open(dest, 'w'), ensure_ascii=False, indent=1)

print('itens:', len(itens))
print('total geral calculado: %.4f' % sum(i['total'] for i in itens))
print('\n--- divergencias item x item (qtd*custo != total):', len(avisos))
for a in avisos[:40]: print('  ', a)
print('\n--- conferencia por bloco ---')
for c in conferencia:
    p = c.get('planilha'); k = c['calculado']
    flag = '' if p is None else ('OK' if abs(p - k) < 0.02 else '<<< DIVERGE')
    print(f"  {c['bloco']:<45} planilha={p} calculado={k} {flag}")
